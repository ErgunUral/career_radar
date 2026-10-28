from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def build_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Ergun Ural")
    document.add_paragraph(
        "QA Manager with 13 yıl experience and 6 yıl management lead experience."
    )
    document.add_paragraph(
        "Fintech ve e-ticaret alanlarında test automation, Jira, Xray, CI/CD, Selenium, Agile ve Scrum."
    )
    document.add_paragraph(
        "ISTQB sertifikası, Computer Engineering eğitimi, English ve Türkçe."
    )
    document.add_paragraph(
        "Remote ve hybrid çalışma tercih edilir. 10 kişilik takım yönetimi."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def create_profile() -> str:
    upload = client.post(
        "/api/v1/files/cv",
        files={
            "file": (
                "ergun-ural-cv.docx",
                build_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"user_id": "jobs-user-1", "language": "tr"},
    )
    file_id = upload.json()["file"]["id"]
    parsed = client.post("/api/v1/cv/parse", json={"file_id": file_id})
    assert parsed.status_code == 200
    return parsed.json()["id"]


def test_import_analyze_and_list_jobs() -> None:
    profile_id = create_profile()
    imported = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-1",
            "profile_id": profile_id,
            "title": "QA Automation Manager",
            "company_name": "Radar Fintech",
            "source_platform": "manual_import",
            "source_url": "https://example.com/jobs/qa-automation-manager",
            "location": "Remote",
            "work_mode": "remote",
            "language": "İngilizce",
            "salary_min": 150000,
            "salary_max": 180000,
            "raw_text": "Fintech şirketi için QA Automation Manager aranıyor. 8 yıl deneyim, 4 yıl liderlik, Python, Jira, Xray, CI/CD, Selenium, Agile ve Scrum deneyimi beklenir.",
        },
    )
    assert imported.status_code == 200
    body = imported.json()
    assert body["analysis_result"]["total_score"] >= 60
    assert body["analysis_result"]["recommendation"] in {
        "Hemen başvur",
        "Yüksek öncelikli",
        "Değerlendirilebilir",
    }

    listed = client.get("/api/v1/jobs", params={"user_id": "jobs-user-1"})
    assert listed.status_code == 200
    assert len(listed.json()) == 1

    detail = client.get(f"/api/v1/jobs/{body['id']}")
    assert detail.status_code == 200
    assert detail.json()["job"]["company_name"] == "Radar Fintech"


def test_create_and_update_application() -> None:
    profile_id = create_profile()
    imported = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-1",
            "profile_id": profile_id,
            "title": "Head of QA",
            "company_name": "Scale Commerce",
            "source_platform": "manual_import",
            "location": "Istanbul",
            "work_mode": "hybrid",
            "raw_text": "E-ticaret şirketi için Head of QA. 10 yıl deneyim, 5 yıl liderlik, Jira, Xray, CI/CD ve test automation deneyimi.",
        },
    )
    job_id = imported.json()["id"]

    application = client.post(
        "/api/v1/applications",
        json={
            "user_id": "jobs-user-1",
            "job_id": job_id,
            "status": "Başvurulacak",
            "notes": "CV özeti güncellenip gönderilecek.",
            "contact_person": "HR Team",
            "expected_salary": 170000,
        },
    )
    assert application.status_code == 200
    application_id = application.json()["id"]

    updated = client.patch(
        f"/api/v1/applications/{application_id}",
        json={
            "status": "Teknik görüşme",
            "notes": "Teknik görüşme 18 Temmuz için planlandı.",
            "offered_salary": 175000,
            "reminder_at": "2026-07-18T09:30:00Z",
        },
    )
    assert updated.status_code == 200
    assert updated.json()["status"] == "Teknik görüşme"
    assert updated.json()["offered_salary"] == 175000
    assert updated.json()["reminder_at"] == "2026-07-18T09:30:00Z"

    listed = client.get("/api/v1/applications", params={"user_id": "jobs-user-1"})
    assert listed.status_code == 200
    assert len(listed.json()) == 1
    assert listed.json()[0]["reminder_at"] == "2026-07-18T09:30:00Z"


def test_list_jobs_filters_active_recent_and_platform() -> None:
    profile_id = create_profile()
    now = datetime.now(timezone.utc)
    recent_iso = now.isoformat()
    older_iso = (now - timedelta(days=45)).isoformat()

    recent = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-filter",
            "profile_id": profile_id,
            "title": "Senior QA Lead",
            "company_name": "Insider",
            "source_platform": "linkedin",
            "source_url": "https://www.linkedin.com/jobs/view/1",
            "location": "Istanbul",
            "work_mode": "remote",
            "raw_text": "Senior QA Lead role for modern SaaS quality organization.",
            "metadata": {"published_at": recent_iso, "is_active": True},
        },
    )
    assert recent.status_code == 200

    inactive = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-filter",
            "profile_id": profile_id,
            "title": "Test Automation Engineer",
            "company_name": "Kariyer Company",
            "source_platform": "kariyernet",
            "source_url": "https://www.kariyer.net/is-ilani/example-1",
            "location": "Istanbul",
            "work_mode": "hybrid",
            "raw_text": "Automation role.",
            "metadata": {"published_at": recent_iso, "is_active": False},
        },
    )
    assert inactive.status_code == 200

    older = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-filter",
            "profile_id": profile_id,
            "title": "QA Manager",
            "company_name": "Legacy Co",
            "source_platform": "company_page",
            "source_url": "https://careers.example.com/jobs/legacy",
            "location": "Ankara",
            "work_mode": "onsite",
            "raw_text": "Legacy QA Manager role.",
            "metadata": {"published_at": older_iso, "is_active": True},
        },
    )
    assert older.status_code == 200

    filtered = client.get(
        "/api/v1/jobs",
        params={
            "user_id": "jobs-user-filter",
            "active_only": "true",
            "published_within_days": 30,
            "source_platform": "linkedin",
            "dedupe": "true",
        },
    )
    assert filtered.status_code == 200
    payload = filtered.json()
    assert len(payload) == 1
    assert payload[0]["company_name"] == "Insider"


def test_list_jobs_excludes_expired_and_closed_applications() -> None:
    profile_id = create_profile()
    now = datetime.now(timezone.utc)

    expired = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-expiry",
            "profile_id": profile_id,
            "title": "Expired QA Lead",
            "company_name": "Closed Corp",
            "source_platform": "company_page",
            "source_url": "https://careers.example.com/jobs/expired-qa-lead",
            "location": "Remote",
            "work_mode": "remote",
            "raw_text": "Expired QA lead role.",
            "metadata": {
                "published_at": now.isoformat(),
                "expires_at": (now - timedelta(days=1)).isoformat(),
                "is_active": True,
            },
        },
    )
    assert expired.status_code == 200

    closed = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-expiry",
            "profile_id": profile_id,
            "title": "Closed Test Manager",
            "company_name": "No More Applications Inc",
            "source_platform": "linkedin",
            "source_url": "https://www.linkedin.com/jobs/view/closed-1",
            "location": "Istanbul",
            "work_mode": "hybrid",
            "raw_text": "Closed test manager role.",
            "metadata": {
                "published_at": now.isoformat(),
                "application_open": False,
            },
        },
    )
    assert closed.status_code == 200

    open_job = client.post(
        "/api/v1/jobs/import",
        json={
            "user_id": "jobs-user-expiry",
            "profile_id": profile_id,
            "title": "Open QA Manager",
            "company_name": "Fresh Jobs Ltd",
            "source_platform": "kariyernet",
            "source_url": "https://www.kariyer.net/is-ilani/open-job-1",
            "location": "Istanbul",
            "work_mode": "remote",
            "raw_text": "Open QA manager role.",
            "metadata": {
                "published_at": now.isoformat(),
                "application_open": True,
                "expires_at": (now + timedelta(days=7)).isoformat(),
            },
        },
    )
    assert open_job.status_code == 200

    filtered = client.get(
        "/api/v1/jobs",
        params={
            "user_id": "jobs-user-expiry",
            "active_only": "true",
            "published_within_days": 30,
            "dedupe": "true",
        },
    )
    assert filtered.status_code == 200
    payload = filtered.json()
    assert len(payload) == 1
    assert payload[0]["company_name"] == "Fresh Jobs Ltd"
